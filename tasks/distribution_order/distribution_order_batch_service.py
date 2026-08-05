# -*- coding: utf-8 -*-
"""
# @Time    : 2026/8/3
# @Author  : Zhu Yaming
# @File    : distribution_order_batch_service.py
# @Description : 一件代发 Excel 批量建单 / 批量审核 / 批量下发
"""
from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from apps.system.distribution_order.constants import (
    CREATE_SOURCE_BATCH_DROPSHIP,
    COOPERATION_METHOD_DROPSHIP,
)
from apps.system.distribution_order.distribution_order_fulfillment_service import (
    _assert_dispatch_stock, _ship_qty, dispatch_order, dispatch_order_test

)
from apps.system.distribution_order.distribution_order_line_import_service import (
    _fetch_valid_product_skus,
)
from apps.system.distribution_order.distribution_order_line_service import (
    _load_alive_lines,
    save_line_details,
)
from apps.system.distribution_order.distribution_order_service import (
    _allocate_batch_id,
    _allocate_order_sn,
    _load_order,
    submit_distribution_order,
)
from apps.system.distribution_order.distribution_order_workflow_engine import (
    ACTION_ORDER_APPROVE,
    OPERATOR_APPROVER,
    STATUS_ORDER_CREATE,
    STATUS_ORDER_REVIEW,
    STATUS_PENDING_DISPATCH,
    TRIGGER_MANUAL,
    apply_transition,
    normalize_line_manager_user_id,
    user_can_review_order,
)
from apps.system.distribution_order.errors import DistributionOrderError
from apps.system.distribution_order.models import (
    DistributionOrder,
    DistributionOrderItemDetail,
    DistributionOrderSnapshot,
)
from apps.system.distribution_order.schemas import (
    BatchApproveIn,
    BatchDispatchIn,
    BatchSubmitIn,
    BatchSubmitRowIn,
    DispatchIn,
    OrderDetailIn,
    OrderSubmitIn,
)
from apps.system.distribution_order.translate import translate_text
from apps.system.offline_customer.models import (
    ContractStatus,
    OfflineContract,
    OfflineCustomer,
)

BATCH_TEMPLATE_HEADERS: Tuple[str, ...] = (
    "客户编码",
    "客户侧订单号",
    "收件人姓名",
    "收件人联系方式",
    "收件人地址-国家",
    "收件人地址-省州",
    "收件人地址-城市",
    "收件人地址-详细地址",
    "收件人地址-邮编",
    "SKU编码",
    "含增值税单价",
    "销售数量",
    "含增值税运费收入",
)

# 前 N 列（至邮编）模板预设为文本格式，避免 Excel 自动改号码/邮编
_BATCH_TEMPLATE_TEXT_COL_COUNT = 9

_FIELD_CN_LABELS: Dict[str, Tuple[str, ...]] = {
    "customer_code": ("客户编码",),
    "customer_po_no": ("客户侧订单号",),
    "name": ("收件人姓名",),
    "contact_info": ("收件人联系方式",),
    "country": ("收件人地址-国家",),
    "province": ("收件人地址-省州",),
    "city": ("收件人地址-城市",),
    "address": ("收件人地址-详细地址",),
    "post_code": ("收件人地址-邮编",),
    "sku": ("SKU编码", "SKU", "sku"),
    "price_with_vat": ("含增值税单价",),
    "qty": ("销售数量", "数量"),
    "freight_with_vat": ("含增值税运费收入",),
}

_MAX_IMPORT_ROWS = 5000


def _msg(
    key: str,
    is_trans: bool = False,
    translation_dict: Optional[Dict[str, Any]] = None,
    **kwargs: Any,
) -> str:
    text = translate_text(key, is_trans, translation_dict or {}) or key
    if kwargs:
        return text.format(**kwargs)
    return text


def build_batch_import_template(
    *,
    is_trans: bool = False,
    translation_dict: Optional[Dict[str, Any]] = None,
) -> Tuple[bytes, str, str]:
    td = translation_dict or {}
    headers = [
        translate_text(h, is_trans, td) or h for h in BATCH_TEMPLATE_HEADERS
    ]
    sheet = (translate_text("批量订单", is_trans, td) or "批量订单")[:31]
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    # 客户/地址类列预设文本格式（与表头前 9 列一致）
    for col_idx in range(1, _BATCH_TEMPLATE_TEXT_COL_COUNT + 1):
        col_letter = get_column_letter(col_idx)
        for row_idx in range(2, _MAX_IMPORT_ROWS + 2):
            ws["{}{}".format(col_letter, row_idx)].number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    ts = datetime.now().strftime("%Y%m%d%H%M")
    display = "一件代发批量导入模板{}.xlsx".format(ts)
    display = translate_text("一件代发批量导入模板", is_trans, td) or "一件代发批量导入模板"
    display = "{}{}.xlsx".format(display, ts)
    ascii_name = "dropship_batch_template_{}.xlsx".format(ts)
    return buf.getvalue(), display, ascii_name


def _normalize_header(name: Any) -> str:
    return str(name or "").strip()


def _normalize_cell_str(val: Any) -> str:
    """Excel 单元格转字符串（文本列；避免号码/邮编被读成数字）。"""
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return ""
    if isinstance(val, bool):
        return str(val)
    if isinstance(val, int):
        return str(val)
    if isinstance(val, float):
        if val == val.to_integral_value():
            return str(int(val))
        s = str(val).strip()
        return "" if s.lower() == "nan" else s
    s = str(val).strip()
    return "" if s.lower() == "nan" else s


_SNAP_COMPARE_KEYS: Tuple[str, ...] = (
    "name", "contact_info", "country", "province", "city", "address", "post_code",
)


def _assert_same_po_recipient(
    po_no: str, base: Dict[str, Any], current: Dict[str, Any],
) -> None:
    """同客户侧订单号：收件人姓名/联系方式/地址须完全一致。"""
    for key in _SNAP_COMPARE_KEYS:
        left = (base.get(key) or "").strip()
        right = (current.get(key) or "").strip()
        if key == "country":
            if left.upper() != right.upper():
                raise DistributionOrderError(
                    40000,
                    "客户侧订单号{}收件人/地址不一致".format(po_no),
                )
            continue
        if left != right:
            raise DistributionOrderError(
                40000,
                "客户侧订单号{}收件人/地址不一致".format(po_no),
            )


def _build_header_map(
    is_trans: bool = False,
    translation_dict: Optional[Dict[str, Any]] = None,
) -> Dict[str, str]:
    mapped: Dict[str, str] = {}
    td = translation_dict or {}
    for field, labels in _FIELD_CN_LABELS.items():
        for label in labels:
            text = _normalize_header(label)
            if text:
                mapped[text.lower()] = field
            trans = _normalize_header(translate_text(label, is_trans, td) or "")
            if trans:
                mapped[trans.lower()] = field
    return mapped


def _parse_decimal(val: Any, *, row_no: int, field: str) -> Decimal:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        raise DistributionOrderError(
            40000, "第{}行{}不能为空".format(row_no, field),
        )
    s = str(val).strip()
    if not s or s.lower() == "nan":
        raise DistributionOrderError(
            40000, "第{}行{}不能为空".format(row_no, field),
        )
    try:
        d = Decimal(s)
    except (InvalidOperation, ValueError):
        raise DistributionOrderError(
            40000, "第{}行{}格式无效".format(row_no, field),
        )
    return d


def _parse_int_qty(val: Any, *, row_no: int) -> int:
    d = _parse_decimal(val, row_no=row_no, field="销售数量")
    if d != d.to_integral_value() or d <= 0:
        raise DistributionOrderError(
            40000, "第{}行销售数量必须为正整数".format(row_no),
        )
    return int(d)


def parse_batch_orders_from_excel(
    content: bytes,
    *,
    expected_customer_code: str,
    expected_customer_country: str = "",
    is_trans: bool = False,
    translation_dict: Optional[Dict[str, Any]] = None,
) -> List[Dict[str, Any]]:
    """解析 Excel，按文件行序返回扁平明细。

    同 PO：SKU 不可重复；收件人姓名/联系方式/地址须一致；国家还须与客户国家一致。
    """
    try:
        df = pd.read_excel(io.BytesIO(content), dtype=object)
    except Exception as exc:
        raise DistributionOrderError(
            40000, "文件解析失败：{}".format(exc),
        ) from exc
    if df is None or df.empty:
        raise DistributionOrderError(40000, "该文件无数据")
    if len(df) > _MAX_IMPORT_ROWS:
        raise DistributionOrderError(
            40000, "上传数据过多，单次最多上传{}条".format(_MAX_IMPORT_ROWS),
        )

    header_map = _build_header_map(is_trans, translation_dict)
    col_to_field: Dict[str, str] = {}
    for col in df.columns:
        key = header_map.get(_normalize_header(col).lower())
        if key:
            col_to_field[col] = key
    required = {
        "customer_code", "customer_po_no", "name", "contact_info", "country",
        "province", "city", "address", "post_code", "sku",
        "price_with_vat", "qty", "freight_with_vat",
    }
    present = set(col_to_field.values())
    missing = required - present
    if missing:
        raise DistributionOrderError(
            40000, "缺少必填列，请使用模板下载的表头",
        )

    expected = (expected_customer_code or "").strip()
    expected_country = (expected_customer_country or "").strip().upper()
    rows: List[Dict[str, Any]] = []
    # 同 PO：国家 + SKU；key=PO
    po_meta: Dict[str, Dict[str, Any]] = {}
    for idx, series in df.iterrows():
        row_no = int(idx) + 2
        raw: Dict[str, Any] = {}
        for col, field in col_to_field.items():
            raw[field] = series.get(col)

        cust_code = _normalize_cell_str(raw.get("customer_code"))
        if not cust_code:
            raise DistributionOrderError(
                40000, "第{}行客户编码不能为空".format(row_no),
            )
        if cust_code != expected:
            raise DistributionOrderError(
                40000,
                "第{}行客户编码与所选客户不一致".format(row_no),
            )
        po_no = _normalize_cell_str(raw.get("customer_po_no"))
        if not po_no:
            raise DistributionOrderError(
                40000, "第{}行客户侧订单号不能为空".format(row_no),
            )
        recipient_name = _normalize_cell_str(raw.get("name"))
        if not recipient_name:
            raise DistributionOrderError(
                40000, "第{}行收件人姓名不能为空".format(row_no),
            )
        sku = _normalize_cell_str(raw.get("sku"))
        if not sku:
            raise DistributionOrderError(
                40000, "第{}行SKU编码不能为空".format(row_no),
            )
        price = _parse_decimal(
            raw.get("price_with_vat"), row_no=row_no, field="含增值税单价",
        )
        if price <= 0:
            raise DistributionOrderError(
                40000, "第{}行含增值税单价必须大于0".format(row_no),
            )
        qty = _parse_int_qty(raw.get("qty"), row_no=row_no)
        freight = _parse_decimal(
            raw.get("freight_with_vat"),
            row_no=row_no, field="含增值税运费收入",
        )
        if freight < 0:
            raise DistributionOrderError(
                40000, "第{}行含增值税运费收入不能为负".format(row_no),
            )

        country = _normalize_cell_str(raw.get("country"))
        if not country:
            raise DistributionOrderError(
                40000, "第{}行收件人地址-国家不能为空".format(row_no),
            )
        if expected_country and country.upper() != expected_country:
            raise DistributionOrderError(
                40000,
                "第{}行收件国家与客户国家不一致".format(row_no),
            )

        snap = {
            "name": recipient_name,
            "contact_info": _normalize_cell_str(raw.get("contact_info")),
            "country": country,
            "province": _normalize_cell_str(raw.get("province")),
            "city": _normalize_cell_str(raw.get("city")),
            "address": _normalize_cell_str(raw.get("address")),
            "post_code": _normalize_cell_str(raw.get("post_code")),
        }
        if po_no not in po_meta:
            po_meta[po_no] = {"snapshot": snap, "skus": {sku}}
        else:
            meta = po_meta[po_no]
            _assert_same_po_recipient(po_no, meta["snapshot"], snap)
            if sku in meta["skus"]:
                raise DistributionOrderError(
                    40000,
                    "客户侧订单号{}内SKU重复：{}".format(po_no, sku),
                )
            meta["skus"].add(sku)

        rows.append({
            "customer_po_no": po_no,
            "name": snap["name"],
            "contact_info": snap["contact_info"],
            "country": snap["country"],
            "province": snap["province"],
            "city": snap["city"],
            "address": snap["address"],
            "post_code": snap["post_code"],
            "sku": sku,
            "price_with_vat": price,
            "qty": qty,
            "freight_with_vat": freight,
        })

    if not rows:
        raise DistributionOrderError(40000, "该文件无有效数据行")
    return rows


async def _load_dropship_customer(
    session: AsyncSession, customer_id: int,
) -> OfflineCustomer:
    row = (
        await session.execute(
            select(OfflineCustomer).where(OfflineCustomer._id == customer_id)
        )
    ).scalar_one_or_none()
    if not row:
        raise DistributionOrderError(40000, "线下客户不存在")
    if int(getattr(row, "cooperation_method", 0) or 0) != COOPERATION_METHOD_DROPSHIP:
        raise DistributionOrderError(40000, "仅一件代发客户支持批量建单")
    return row


async def _load_effective_contract(
    session: AsyncSession,
    customer: OfflineCustomer,
    contract_id: Optional[int],
) -> Optional[OfflineContract]:
    cid = contract_id or getattr(customer, "current_contract_id", None)
    if not cid:
        return None
    row = (
        await session.execute(
            select(OfflineContract).where(OfflineContract._id == int(cid))
        )
    ).scalar_one_or_none()
    if row and int(row.status) == ContractStatus.EFFECTIVE:
        return row
    return row


async def _assert_skus_enabled(
    session: AsyncSession, skus: List[str],
) -> None:
    if not skus:
        raise DistributionOrderError(40000, "订单明细不能为空")
    valid_skus = await _fetch_valid_product_skus(session, skus)
    invalid = sorted({
        s for s in skus if str(s).strip().upper() not in valid_skus
    })
    if invalid:
        raise DistributionOrderError(
            40000,
            "以下SKU不存在或未启用：{}".format("、".join(invalid[:20])),
        )


def _serialize_batch_parse_orders(
    rows: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    """按 Excel 行序返回扁平记录，并预留红线/指导价/库存字段供前端填充。"""
    out: List[Dict[str, Any]] = []
    for row in rows:
        item = dict(row)
        item["guide_price"] = None
        item["red_line_price"] = None
        item["stock_on_hand"] = None
        item["stock_in_transit"] = None
        item["stock_planned"] = None
        out.append(item)
    return out


async def parse_batch_dropship_orders(
    session: AsyncSession,
    *,
    offline_customer_id: int,
    file_content: bytes,
    is_trans: bool = False,
    translation_dict: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """解析 Excel + SKU 校验，不落库。前端匹配红线价/库存后再调 submit。"""
    customer = await _load_dropship_customer(session, offline_customer_id)
    parsed_rows = parse_batch_orders_from_excel(
        file_content,
        expected_customer_code=(customer.customer_code or "").strip(),
        expected_customer_country=(customer.customer_country or "").strip(),
        is_trans=is_trans,
        translation_dict=translation_dict,
    )
    all_skus = [r["sku"] for r in parsed_rows]
    await _assert_skus_enabled(session, all_skus)
    rows = _serialize_batch_parse_orders(parsed_rows)
    return {
        "offline_customer_id": int(customer._id),
        "customer_code": customer.customer_code or "",
        "order_details": rows,
        "imported_count": len(rows),
        "order_count": len({r["customer_po_no"] for r in rows}),
    }


def _group_submit_rows(
    rows: List[BatchSubmitRowIn],
    *,
    expected_customer_country: str = "",
) -> List[Dict[str, Any]]:
    """按客户侧订单号聚合；组顺序=该 PO 首次出现顺序，行顺序=提交列表顺序。

    同 PO：收件人姓名/联系方式/地址须一致，国家还须与客户国家一致；SKU 不重复。
    """
    expected_country = (expected_customer_country or "").strip().upper()
    groups: Dict[str, Dict[str, Any]] = {}
    for row in rows:
        po = (row.customer_po_no or "").strip()
        if not po:
            raise DistributionOrderError(40000, "客户侧订单号不能为空")
        sku = (row.sku or "").strip()
        if not sku:
            raise DistributionOrderError(40000, "SKU编码不能为空")
        if row.price_with_vat is None or row.price_with_vat <= 0:
            raise DistributionOrderError(40000, "含增值税单价必须大于0")
        if row.freight_with_vat is None or row.freight_with_vat < 0:
            raise DistributionOrderError(40000, "含增值税运费收入不能为负")

        country = (row.country or "").strip()
        if not country:
            raise DistributionOrderError(40000, "收件人地址-国家不能为空")
        if expected_country and country.upper() != expected_country:
            raise DistributionOrderError(40000, "收件国家与客户国家不一致")
        recipient_name = (row.name or "").strip()
        if not recipient_name:
            raise DistributionOrderError(40000, "收件人姓名不能为空")

        snap = {
            "name": recipient_name,
            "contact_info": (row.contact_info or "").strip(),
            "country": country,
            "province": (row.province or "").strip(),
            "city": (row.city or "").strip(),
            "address": (row.address or "").strip(),
            "post_code": (row.post_code or "").strip(),
        }
        line = OrderDetailIn(
            sku=sku,
            price_with_vat=row.price_with_vat,
            qty=row.qty,
            freight_with_vat=row.freight_with_vat,
            guide_price=row.guide_price,
            red_line_price=row.red_line_price,
            stock_on_hand=row.stock_on_hand,
            stock_in_transit=row.stock_in_transit,
            stock_planned=row.stock_planned,
        )
        if po not in groups:
            groups[po] = {
                "customer_po_no": po,
                "snapshot": snap,
                "lines": [line],
                "skus": {sku},
            }
            continue
        g = groups[po]
        _assert_same_po_recipient(po, g["snapshot"], snap)
        if sku in g["skus"]:
            raise DistributionOrderError(
                40000,
                "客户侧订单号{}内SKU重复：{}".format(po, sku),
            )
        g["skus"].add(sku)
        g["lines"].append(line)
    if not groups:
        raise DistributionOrderError(40000, "订单明细不能为空")
    return list(groups.values())


async def submit_batch_dropship_orders(
    session: AsyncSession,
    body: BatchSubmitIn,
    *,
    operator_id: Optional[int] = None,
    exchange_rate_dict: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """按前端提交的扁平行落库并直进订单审核(40)。整批同事务。

    注：框架报价校验下期再做；本期 SKU 校验商品库已启用记录。
    """
    if not body.order_details:
        raise DistributionOrderError(40000, "订单明细不能为空")

    lm_user_id = normalize_line_manager_user_id(body.order_lm_user_id)

    customer = await _load_dropship_customer(session, body.offline_customer_id)
    groups = _group_submit_rows(
        body.order_details,
        expected_customer_country=(customer.customer_country or "").strip(),
    )
    all_skus = [ln.sku for g in groups for ln in g["lines"]]
    await _assert_skus_enabled(session, all_skus)

    contract = await _load_effective_contract(
        session, customer, body.contract_id,
    )
    country = (customer.customer_country or "").strip()
    if not country:
        raise DistributionOrderError(40000, "客户国家为空，无法生成单号")

    batch_id = await _allocate_batch_id(session)
    order_ids: List[int] = []
    order_sns: List[str] = []

    for group in groups:
        order_sn = await _allocate_order_sn(session, country)
        order = DistributionOrder(
            order_sn=order_sn,
            status=STATUS_ORDER_CREATE,
            offline_customer_id=int(customer._id),
            customer_code=customer.customer_code or "",
            customer_short_name=customer.customer_short_name or "",
            customer_country=customer.customer_country,
            customer_type_first=customer.customer_type_first,
            customer_type_second=customer.customer_type_second,
            cooperation_method=getattr(customer, "cooperation_method", None),
            batch_id=batch_id,
            customer_po_no=group["customer_po_no"],
            create_source=CREATE_SOURCE_BATCH_DROPSHIP,
            create_by=operator_id,
            update_by=operator_id,
            sales_user_id=body.sales_user_id,
            shop_id=(body.shop_id or "").strip() or None,
            is_tax_free=int(body.is_tax_free or 0),
            is_ewt=int(body.is_ewt or 0),
            vat_rate=body.vat_rate,
            ewt_rate=body.ewt_rate,
            currency=body.currency or customer.settlement_currency,
            order_delivery_fee_payment=body.order_delivery_fee_payment,
            ship_method=body.ship_method,
            expected_ship_date=body.expected_ship_date,
            remark=body.remark,
            order_lm_user_id=lm_user_id,
        )
        if contract is not None:
            order.contract_id = int(contract._id)
            order.is_prepayment = int(contract.is_prepayment or 0)
            order.prepayment_ratio = contract.prepayment_ratio or 0
            order.settlement_method = contract.settlement_method
            if body.order_delivery_fee_payment is None:
                order.order_delivery_fee_payment = contract.delivery_fee_payment
        session.add(order)
        await session.flush()

        snap_data = group["snapshot"]
        snap = DistributionOrderSnapshot(order_id=int(order._id))
        snap.name = snap_data.get("name")
        snap.contact_info = snap_data.get("contact_info")
        snap.country = snap_data.get("country")
        snap.province = snap_data.get("province")
        snap.city = snap_data.get("city")
        snap.address = snap_data.get("address")
        snap.post_code = snap_data.get("post_code")
        if contract is not None:
            snap.contract_no = contract.contract_no
            snap.effective_start = contract.effective_start
            snap.effective_end = contract.effective_end
            snap.settlement_days = contract.settlement_days
            snap.settlement_currency = customer.settlement_currency
            snap.delivery_method = contract.delivery_method
            snap.carrier = contract.carrier
            snap.delivery_fee_payment = contract.delivery_fee_payment
            snap.sample_policy = contract.sample_policy
            snap.sample_discount = contract.sample_discount
            snap.bank_account_confirmation = contract.bank_account_confirmation
            snap.mt_delivery_mode = contract.mt_delivery_mode
            snap.owner_staff_id = contract.owner_staff_id
            try:
                snap.effective_node = (
                    int(contract.effective_node)
                    if contract.effective_node not in (None, "")
                    else None
                )
            except (TypeError, ValueError):
                snap.effective_node = None
        session.add(snap)

        await save_line_details(
            session, int(order._id), "order", group["lines"],
            check_status=False,
            replace_all=True,
        )

        await submit_distribution_order(
            session,
            OrderSubmitIn(
                order_id=int(order._id),
                chain_code=body.chain_code,
                order_lm_user_id=lm_user_id,
            ),
            operator_id=operator_id,
            exchange_rate_dict=exchange_rate_dict,
        )
        order_ids.append(int(order._id))
        order_sns.append(order_sn)

    return {
        "batch_id": batch_id,
        "order_ids": order_ids,
        "order_sns": order_sns,
        "count": len(order_ids),
    }


async def batch_approve_orders(
    session: AsyncSession,
    body: BatchApproveIn,
    *,
    operator_id: Optional[int] = None,
) -> Dict[str, Any]:
    """批量审核通过；驳回请走单条 approval/order/reject（每单 remark）。"""
    success: List[Dict[str, Any]] = []
    failed: List[Dict[str, Any]] = []
    remark = (body.remark or "").strip() or None

    for oid in body.order_ids:
        try:
            order = await _load_order(session, int(oid))
            if int(getattr(order, "create_source", 0) or 0) != CREATE_SOURCE_BATCH_DROPSHIP:
                raise DistributionOrderError(40000, "仅批量创建订单支持批量审核")
            if int(order.status) != STATUS_ORDER_REVIEW:
                raise DistributionOrderError(40000, "订单未到订单审核状态")
            if not await user_can_review_order(session, order, operator_id):
                raise DistributionOrderError(40300, "您暂无审批该订单的权限")
            _, meta, _msg = await apply_transition(
                session,
                order_id=int(oid),
                action=ACTION_ORDER_APPROVE,
                operator_id=operator_id,
                operator_role=OPERATOR_APPROVER,
                trigger_type=TRIGGER_MANUAL,
                remark=remark,
            )
            success.append({
                "order_id": int(oid),
                "order_sn": order.order_sn,
                "to_status": meta.get("to_status"),
            })
        except DistributionOrderError as exc:
            failed.append({"order_id": int(oid), "msg": exc.msg})
        except Exception as exc:
            failed.append({"order_id": int(oid), "msg": str(exc)})

    return {
        "success": success,
        "failed": failed,
        "success_count": len(success),
        "failed_count": len(failed),
    }


async def batch_dispatch_orders(
    session: AsyncSession,
    body: BatchDispatchIn,
    *,
    operator_id: Optional[int] = None,
    operator_name: Optional[str] = None,
    shop_dict: Optional[Dict[Any, Any]] = None,
) -> Dict[str, Any]:
    success: List[Dict[str, Any]] = []
    skipped: List[Dict[str, Any]] = []
    failed: List[Dict[str, Any]] = []

    for oid in body.order_ids:
        try:
            order = await _load_order(session, int(oid))
            if int(order.status) != STATUS_PENDING_DISPATCH:
                raise DistributionOrderError(40000, "当前状态不允许下发，仅待下发可操作")
            lines = list(
                await _load_alive_lines(
                    session, int(oid), DistributionOrderItemDetail,
                )
            )
            shippable = [ln for ln in lines if _ship_qty(ln) > 0]
            if not shippable:
                raise DistributionOrderError(40000, "无可下发数量的订单明细")
            try:
                _assert_dispatch_stock(shippable)
            except DistributionOrderError as stock_exc:
                skipped.append({
                    "order_id": int(oid),
                    "order_sn": order.order_sn,
                    "msg": stock_exc.msg,
                })
                continue
            data = await dispatch_order_test(
                session,
                DispatchIn(order_id=int(oid)),
                operator_id=operator_id,
                operator_name=operator_name,
                shop_dict=shop_dict,
            )
            success.append(data)
        except DistributionOrderError as exc:
            failed.append({"order_id": int(oid), "msg": exc.msg})
        except Exception as exc:
            failed.append({"order_id": int(oid), "msg": str(exc)})

    return {
        "success": success,
        "skipped": skipped,
        "failed": failed,
        "success_count": len(success),
        "skipped_count": len(skipped),
        "failed_count": len(failed),
    }
